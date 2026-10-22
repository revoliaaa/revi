"""Build the DOCX version of the FPIC guideline from its Markdown source.

Usage: python3 build/build_docx.py   (run from fpic-hd-pasir/)
Requires pandoc and python-docx.
"""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "Pedoman_FPIC_PlanVivo_HD_Pasir.md"
OUT = ROOT / "Pedoman_FPIC_PlanVivo_HD_Pasir.docx"
ACCENT = RGBColor(0x1F, 0x4E, 0x3D)


def set_borders(table):
    """Horizontal rules only: no vertical lines, no boxed cells."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, val in (("top", "single"), ("bottom", "single"), ("insideH", "single"),
                      ("left", "nil"), ("right", "nil"), ("insideV", "nil")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        if val == "single":
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "9AA5A0")
        borders.append(el)
    tbl_pr.append(borders)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def main():
    subprocess.run(["pandoc", str(SRC), "-o", str(OUT), "--toc", "--toc-depth=2"], check=True)
    doc = Document(str(OUT))

    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21.0)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2))

    styles = doc.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in styles:
            st = styles[name]
            st.font.name = "Georgia"
            st.font.size = Pt(10)
            st.paragraph_format.line_spacing = 1.3
    for name, size in (("Title", 22), ("Subtitle", 13), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        if name in styles:
            st = styles[name]
            st.font.name = "Georgia"
            st.font.size = Pt(size)
            st.font.color.rgb = ACCENT
            st.paragraph_format.space_before = Pt(14 if name.startswith("Heading") else 0)
            st.paragraph_format.space_after = Pt(4)

    for table in doc.tables:
        set_borders(table)
        for r_idx, row in enumerate(table.rows):
            for cell in row.cells:
                if r_idx == 0:
                    shade(cell, "E4ECE8")
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.1
                    for run in p.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(8)
                        if r_idx == 0:
                            run.font.bold = True

    doc.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
