# -*- coding: utf-8 -*-
"""Satu slide presentasi: 8 tahap Hutan Desa Pasir, dari persiapan PIN sampai penjualan.

    python3 build_one_slide.py ../Roadmap_1Slide_HD_Pasir.pptx
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from build_pptx import (AMBER, DISPLAY, HAIR, INK, MUTED, NAVY, NAVY_MID, NAVY_TINT, TEAL, TEAL_MID, TEAL_TINT,
                        WHITE, box, line, rich, shape_text, text)
from data import TAHAP

AMBER_TINT = RGBColor(0xFB, 0xF1, 0xE4)
GREY_TINT = RGBColor(0xEE, 0xEF, 0xEC)

# Isi kartu per tahap (judul, periode & milestone diambil dari data.TAHAP).
CARD = {
    1: (["Safeguard: FPIC, GRM, E&S, GESI", "Kebutuhan & progres berjalan",
         "Baseline (PV tools) → project scenario"], "Bahan PIN lengkap", False),
    2: (["PIN Template V1.4", "Submit + PIN review fee", "Review CAR/NIR, maks. 3 putaran"], "PIN disetujui", True),
    3: (["PIN tayang di pipeline Plan Vivo", "Profil proyek & penjajakan investor",
         "Pendamping teregistrasi"], "Proyek di Pipeline", False),
    4: (["Pre-financing investor / offtaker", "Metodologi disetujui Plan Vivo", "Stok karbon, baseline, PDD review"],
        "PDD lolos review", False),
    5: (["DPP tercatat di Kemenhut", "Audit VVB / microscale IE", "MoU & PV Climate Registry"],
        "Proyek terdaftar", True),
    6: (["Perlindungan & restorasi", "Livelihood KUPS & WFD", "Annual Report tiap 12 bln"],
        "Data monitoring", False),
    7: (["Audit verifikasi ≤5 th", "Persetujuan Menteri", "vPVC terbit, nomor seri"], "vPVC di registry", True),
    8: (["Realisasi offtake & penjualan", "Otorisasi/CA bila ekspor; PNBP", "≥60% pendapatan ke komunitas"],
        "Pendapatan & manfaat", False),
}


# Label ringkas agar muat di kartu (nama lengkap tetap di data.TAHAP).
SHORT_TITLE = {3: "Masuk\nPipeline PV"}


def zone_color(no):
    return TEAL if no <= 3 else NAVY


def build(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    # latar
    bg = box(s, 0, 0, 13.333, 7.5, fill=RGBColor(0xFB, 0xFB, 0xF9))

    text(s, 0.45, 0.3, 11, 0.25, "COMMUNITY CARBON HUTAN DESA PASIR · MEMPAWAH, KALIMANTAN BARAT · PLAN VIVO PV CLIMATE",
         size=9, bold=True, color=TEAL)
    text(s, 0.45, 0.55, 12.4, 0.85,
         "SETAPAK 4 berakhir dengan proyek masuk Pipeline Plan Vivo — status ini\n"
         "membuka investasi untuk PDD, validasi, hingga penjualan kredit karbon",
         size=21, font=DISPLAY, color=INK)

    # grid kolom
    lab_x, lab_w = 0.45, 0.85
    x0, x1, gap = 1.3, 12.93, 0.06
    w = (x1 - x0 - 7 * gap) / 8
    cx = [x0 + i * (w + gap) for i in range(8)]

    # pita zona
    zy = 1.62
    zones = [(0, 2, TEAL, "STEP 1 · PROGRAM SETAPAK 4 (TAF)"),
             (3, 7, NAVY, "STEP 2 · PENGEMBANGAN LANJUTAN (PASCA-SETAPAK 4) →")]
    for a, z, col, lbl in zones:
        zb = box(s, cx[a], zy, cx[z] + w - cx[a], 0.32, fill=col)
        shape_text(zb, lbl, size=8.5)

    # garis alur & penanda
    ny = 2.3
    line(s, cx[0] + w / 2, ny, cx[3] - gap / 2, ny, color=TEAL, w=2.5)
    line(s, cx[3] - gap / 2, ny, cx[7] + w / 2, ny, color=NAVY, w=2.5)
    text(s, lab_x, ny - 0.12, lab_w, 0.25, "TAHAP", size=8, bold=True, color=MUTED)

    # batas SETAPAK 4: setelah Tahap 3 (masuk pipeline)
    gx = cx[3] - gap / 2
    line(s, gx, zy - 0.06, gx, 6.0, color=AMBER, w=1.75, dash=MSO_LINE.DASH)
    text(s, gx - 1.6, zy - 0.27, 3.2, 0.22, "▼ GERBANG PIPELINE · AKHIR SETAPAK 4", size=7.5, bold=True, color=AMBER,
         align=PP_ALIGN.CENTER)

    T = {t["no"]: t for t in TAHAP}
    cy, ch = 2.68, 2.4
    for i in range(8):
        no = i + 1
        t = T[no]
        col = zone_color(no)
        bullets, output, ext = CARD[no]
        mx = cx[i] + w / 2
        if no == 3:
            box(s, mx - 0.34, ny - 0.34, 0.68, 0.68, fill=WHITE, line=TEAL, shape=MSO_SHAPE.OVAL, lw=2)
            d = box(s, mx - 0.26, ny - 0.26, 0.52, 0.52, fill=TEAL, shape=MSO_SHAPE.OVAL)
            shape_text(d, "3", size=14)
        elif ext:
            d = box(s, mx - 0.27, ny - 0.27, 0.54, 0.54, fill=WHITE, line=col, shape=MSO_SHAPE.DIAMOND, lw=2.25)
            shape_text(d, str(no), size=13, color=col)
        else:
            d = box(s, mx - 0.24, ny - 0.24, 0.48, 0.48, fill=col, shape=MSO_SHAPE.OVAL)
            shape_text(d, str(no), size=13)

        tint = TEAL_TINT if no <= 3 else NAVY_TINT
        key = no == 3
        card = box(s, cx[i], cy, w, ch, fill=TEAL_TINT if key else WHITE, line=TEAL if key else HAIR,
                   lw=1.75 if key else 0.75)
        box(s, cx[i], cy, w, 0.05, fill=col)
        if key:
            badge = box(s, cx[i] + 0.12, cy - 0.13, w - 0.24, 0.24, fill=TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            shape_text(badge, "★ OUTPUT UTAMA", size=7.5)
        text(s, cx[i] + 0.08, cy + 0.1, w - 0.16, 0.45, SHORT_TITLE.get(no, t["title"]), size=9.5, bold=True, color=INK)
        text(s, cx[i] + 0.08, cy + 0.56, w - 0.16, 0.2, t["ms"], size=6.5,
             bold=True, color=col)
        rich(s, cx[i] + 0.04, cy + 0.8, w - 0.1, 1.1, [[("• ", True, col), (b, False, INK)] for b in bullets],
             size=7.5)
        ob = box(s, cx[i] + 0.08, cy + ch - 0.42, w - 0.16, 0.32, fill=TEAL if key else tint,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        shape_text(ob, "→ " + output, size=7, bold=True, color=WHITE if key else col)

    # lajur dukungan
    sy = 5.22
    text(s, lab_x, sy + 0.04, lab_w, 0.3, "DUKUNGAN", size=8, bold=True, color=MUTED)
    for a, z, fill, fg, lbl in [
        (0, 2, TEAL_TINT, TEAL, "TAF · SETAPAK 4 — SAMPAN Kalimantan & LPHD Pasir"),
        (3, 7, NAVY_TINT, NAVY, "Investor / offtaker (pre-financing, presale berjangka) + pendanaan lanjutan  →  pendapatan karbon"),
    ]:
        b = box(s, cx[a], sy, cx[z] + w - cx[a], 0.32, fill=fill)
        shape_text(b, lbl, size=8, bold=True, color=fg)

    # lajur unit karbon
    uy = 5.62
    text(s, lab_x, uy + 0.04, lab_w, 0.3, "UNIT KARBON", size=8, bold=True, color=MUTED)
    b = box(s, cx[0], uy, cx[3] + w - cx[0], 0.32, fill=GREY_TINT)
    shape_text(b, "Belum ada kredit — pipeline ≠ registrasi ≠ issuance", size=8, bold=False, color=MUTED)
    for i, lbl, fill, fg in [(4, "fPVC (restorasi)", NAVY_TINT, NAVY), (5, "rPVC (Annual Report)", NAVY_TINT, NAVY),
                             (6, "vPVC ✓ dapat di-retire", NAVY, WHITE), (7, "Dijual / diserahkan", NAVY, WHITE)]:
        b = box(s, cx[i], uy, w, 0.32, fill=fill)
        shape_text(b, lbl, size=7.5, bold=True, color=fg)

    # tanggal kunci
    ky = 6.2
    line(s, 0.45, ky - 0.08, 12.9, ky - 0.08, color=HAIR, w=0.75)
    keys = [("Pipeline PV", "Output utama SETAPAK 4", TEAL),
            ("≤ 3 tahun", "PDD wajib diajukan sejak PIN disetujui", NAVY),
            ("≤ 5 tahun", "Verifikasi pertama setelah validasi", NAVY),
            ("≥ 60%", "pendapatan PVC untuk komunitas", NAVY)]
    kw = 2.55
    for i, (big, cap, col) in enumerate(keys):
        kx = 0.45 + i * (kw + 0.12)
        text(s, kx, ky, kw, 0.42, big, size=20, bold=True, font=DISPLAY, color=col)
        text(s, kx, ky + 0.45, kw, 0.3, cap, size=8, color=INK)
    # legenda
    lx = 11.2
    o = box(s, lx, ky + 0.05, 0.2, 0.2, fill=MUTED, shape=MSO_SHAPE.OVAL)
    text(s, lx + 0.27, ky + 0.02, 1.5, 0.25, "Dikerjakan tim", size=8, color=MUTED)
    d = box(s, lx - 0.02, ky + 0.36, 0.24, 0.24, fill=WHITE, line=MUTED, shape=MSO_SHAPE.DIAMOND, lw=1.5)
    text(s, lx + 0.27, ky + 0.35, 1.6, 0.25, "Keputusan pihak luar", size=8, color=MUTED)
    line(s, lx - 0.02, ky + 0.78, lx + 0.22, ky + 0.78, color=AMBER, w=1.75, dash=MSO_LINE.DASH)
    text(s, lx + 0.27, ky + 0.66, 1.6, 0.25, "Batas SETAPAK 4", size=8, color=MUTED)

    text(s, 0.45, 7.17, 12.4, 0.22,
         "Sumber: Plan Vivo PV Climate Procedures Manual V3.9 & Project Requirements V5.7; Permenhut 6/2026. "
         "Keputusan Plan Vivo, VVB dan Menteri berada di luar kendali proyek. Presale = kontrak berjangka; unit diserahkan setelah terbit.",
         size=7, color=MUTED)

    s.notes_slide.notes_text_frame.text = (
        "Alur: Tahap 1–3 adalah cakupan program SETAPAK 4 yang didanai TAF. Output utamanya "
        "Tahap 3: PIN disetujui dan proyek masuk pipeline Plan Vivo. Pipeline belum berarti registrasi atau kredit. "
        "Garis gerbang setelah Tahap 3 menandai akhir SETAPAK 4. Dengan status pipeline, proyek mencari investor atau "
        "pembeli untuk pre-financing/offtake berjangka yang membiayai PDD dan validasi (Tahap 4 dst.). Kredit baru muncul setelah registrasi (fPVC untuk restorasi, rPVC dari Annual Report) dan yang "
        "dapat di-retire hanya vPVC setelah verifikasi. Minimal 60% pendapatan dari penjualan untuk peserta dan pihak lokal.")
    return s


def main(out):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    build(prs)
    prs.save(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "one_slide.pptx")
