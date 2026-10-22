# -*- coding: utf-8 -*-
"""Bangun Project Monitoring Excel Hutan Desa Pasir.

    python3 build_xlsx.py ../Milestone_Matrix_HD_Pasir_PVClimate.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from data import (ACTIVITIES, BASIS, COMPLETION, DOC_CHECKLIST, MILESTONES, MS_TAHAP, PLOT_MONTHS, RISKS, ROLES,
                  SOURCES, STEP2, TAF_END, TAHAP, TOOLS, TRANSITION, month_label)

TEAL, TEAL_T, TEAL_M = "0E7C6B", "E3F1EE", "7FBCB1"
NAVY, NAVY_T = "1F2A44", "E6E9EF"
AMBER, AMBER_T = "C26A0A", "F6EBDD"
HAIR = "D5DBD6"

thin = Side(style="thin", color=HAIR)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS = ["Belum mulai", "Berjalan", "Selesai", "Tertunda", "Tidak berlaku"]


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def header_row(ws, row, headers, color):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill(color)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, t, sub):
    ws["A1"] = t
    ws["A1"].font = Font(bold=True, size=14, color=TEAL)
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, size=9, color="5F6B65")


def status_validation(ws, col, first, last):
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}{first}:{col}{last}")
    ws.conditional_formatting.add(f"{col}{first}:{col}{last}",
                                  CellIsRule(operator="equal", formula=['"Selesai"'], fill=fill(TEAL_T)))
    ws.conditional_formatting.add(f"{col}{first}:{col}{last}",
                                  CellIsRule(operator="equal", formula=['"Tertunda"'], fill=fill(AMBER_T)))


def tahap_label(ms):
    t = TAHAP[MS_TAHAP[ms] - 1]
    return f"Tahap {t['no']} — {t['title']}"


def role(k):
    return ROLES.get(k, k)


def sheet_readme(wb):
    ws = wb.active
    ws.title = "README"
    title(ws, "Integrated Community Carbon Development Roadmap – Hutan Desa Pasir",
          "Plan Vivo PV Climate · STEP 1 = SETAPAK 4 (M1–M7, Okt 2026–Sep 2027) · STEP 2 = Post-SETAPAK 4 (M9–M13) · Bulan 1 = Oktober 2026")
    rows = [
        ("Batas SETAPAK 4", "Target program = PIN diajukan (M6) DAN PIN disetujui sehingga proyek tercantum di Plan Vivo Project pipeline (M7). "
                            "STEP 2 ditampilkan untuk perencanaan, bukan kewajiban capaian SETAPAK 4."),
        ("Pipeline ≠ registrasi", "Pipeline = PIN disetujui. Registrasi hanya setelah validasi (VVB atau proses microscale) + MoU Plan Vivo. "
                                 "Issuance PVC lewat Annual Report (fPVC/rPVC) dan verifikasi (vPVC)."),
        ("Tanggal eksternal", "Semua tanggal yang bergantung pada Plan Vivo, VVB/IE atau pemerintah bersifat indikatif."),
        ("Jenis", "INT = milestone internal; EXT = persetujuan eksternal"),
    ] + [(f"Basis {k}", v) for k, v in BASIS.items()] + [("Peran " + k, v) for k, v in ROLES.items()]
    for i, (k, v) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v).alignment = WRAP
    widths(ws, [24, 120])


def sheet_matrix(wb):
    ws = wb.create_sheet("Master_Matrix")
    title(ws, "B. Master Milestone Matrix", "STEP 1 rinci sampai sub-aktivitas; STEP 2 pada tingkat milestone & workstream")
    headers = ["Step", "Programme", "Tahap", "Milestone", "Kode", "Activities", "Sub-activities", "Key Deliverables",
               "Bukti verifikasi", "PIC (Responsible)", "Accountable", "Dependencies", "Completion Criteria",
               "Mulai", "Selesai", "Durasi (bln)", "Jenis", "Basis", "Participatory tools", "Jalur kritis",
               "Status", "% progres", "Catatan"]
    header_row(ws, 4, headers, NAVY)
    r = 5
    titles = {m["code"]: m for m in MILESTONES}
    for m in [m for m in MILESTONES if m["step"] == 1]:
        vals = ["STEP 1", "SETAPAK 4 (TAF)", tahap_label(m["code"]), f"{m['code']} {m['title']}", m["code"], "", "", m["output"].replace("\n", " "),
                "", "", "", "", COMPLETION[m["code"]], month_label(m["start"]), month_label(m["end"]),
                m["end"] - m["start"] + 1, m["type"], "", "", "", "Belum mulai", 0, ""]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(bold=True, color=TEAL)
            c.fill = fill(TEAL_T)
            c.alignment = WRAP
            c.border = BORDER
        r += 1
        for a in [a for a in ACTIVITIES if a[1] == m["code"]]:
            code, ms, name, sub, deliv, evid, pic, acc, dep, st, en, basis, tools, crit = a
            vals = ["STEP 1", "SETAPAK 4 (TAF)", tahap_label(ms), ms, code, name, sub, deliv, evid, role(pic), role(acc), dep,
                    "", month_label(st), month_label(en), en - st + 1, "EXT" if ms == "M7" and code == "M7.5" else "INT",
                    basis, tools, "Ya" if crit else "", "Belum mulai", 0, ""]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=i, value=v)
                c.alignment = WRAP
                c.border = BORDER
            r += 1
    for s2 in STEP2:
        m = titles[s2["code"]]
        vals = ["STEP 2", "Post-SETAPAK 4 (pendanaan lanjutan)", tahap_label(m["code"]), f"{m['code']} {m['title']}", m["code"],
                s2["workstreams"], "", m["output"].replace("\n", " "), "", role("PM") + " / Project Coordinator",
                role("DIR"), s2["prereq"], COMPLETION[m["code"]], m["years"], "", "", m["type"], "PV/ID",
                "", "", "Belum mulai", 0, s2["deadline"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP
            c.border = BORDER
            c.fill = fill(NAVY_T)
        ws.cell(row=r, column=4).font = Font(bold=True, color=NAVY)
        r += 1
    status_validation(ws, "U", 5, r - 1)
    widths(ws, [8, 16, 22, 26, 7, 30, 48, 28, 26, 22, 22, 22, 40, 10, 10, 8, 7, 8, 16, 8, 13, 9, 30])
    ws.auto_filter.ref = f"A4:W{r - 1}"


def sheet_gantt(wb):
    ws = wb.create_sheet("Gantt_Step1")
    title(ws, "C5/C6. Timeline STEP 1 & critical path", "Hijau tua = jalur kritis · hijau muda = ada float · oranye = bergantung review Plan Vivo · header oranye = setelah TAF berakhir (Sep 2027), buffer review")
    months = PLOT_MONTHS
    headers = ["Kode", "Aktivitas", "PIC"] + [month_label(m) for m in range(1, months + 1)]
    header_row(ws, 4, headers, TEAL)
    for m in range(1, months + 1):
        if m > TAF_END:
            ws.cell(row=4, column=3 + m).fill = fill(AMBER)
    r = 5
    for a in ACTIVITIES:
        code, ms, name, *_ = a
        pic, st, en, crit = a[6], a[9], a[10], a[13]
        ws.cell(row=r, column=1, value=code).font = Font(bold=crit, color=TEAL if crit else "5F6B65")
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=pic)
        for m in range(1, months + 1):
            c = ws.cell(row=r, column=3 + m)
            c.border = BORDER
            if st <= m <= en:
                c.fill = fill(AMBER if ms == "M7" else (TEAL if crit else TEAL_M))
        r += 1
    widths(ws, [7, 44, 7] + [7.5] * months)


def simple_sheet(wb, name, ttl, sub, headers, rows, w, color=TEAL, status_col=None):
    ws = wb.create_sheet(name)
    title(ws, ttl, sub)
    header_row(ws, 4, headers, color)
    for r, row in enumerate(rows, 5):
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP
            c.border = BORDER
    if status_col:
        status_validation(ws, status_col, 5, 4 + len(rows))
    widths(ws, w)
    return ws


def main(out):
    wb = Workbook()
    sheet_readme(wb)
    sheet_matrix(wb)
    sheet_gantt(wb)
    simple_sheet(wb, "Tahapan", "Kerangka 8 tahap (Tahap → Milestone)",
                 "Tahap 1–3 = SETAPAK 4 (Tahap 3 = output utama: masuk Pipeline PV) · Tahap 4–8 = pengembangan lanjutan",
                 ["Tahap", "Nama", "Milestone", "Periode (indikatif)", "Cakupan", "Isi utama"],
                 [(t["no"], t["title"], t["ms"], t["period"],
                   {1: "SETAPAK 4", 0: "Transisi", 2: "Post-SETAPAK 4"}[t["step"]], "; ".join(t["points"]))
                  for t in TAHAP], [7, 40, 10, 20, 16, 90])
    simple_sheet(wb, "Participatory_Tools", "C7. Integrasi 17 Plan Vivo Participatory Tools",
                 "Status PIN/PDD mengikuti tabel 'Tools in the Plan Vivo process' (Participatory Toolkit V1.0)",
                 ["Tool", "Nama", "Dipakai di aktivitas", "Status tahap PIN", "Status tahap PDD"],
                 TOOLS, [6, 38, 34, 30, 34])
    simple_sheet(wb, "Doc_Checklist", "C8. Document readiness checklist (STEP 1)", "Isi kolom Status & Lokasi file",
                 ["ID", "Dokumen", "Sumber (aktivitas)", "Basis", "Status", "Lokasi file / tautan", "Catatan"],
                 [d + ("Belum mulai", "", "") for d in DOC_CHECKLIST], [6, 58, 14, 8, 14, 30, 30], status_col="E")
    simple_sheet(wb, "Risk_Register", "C9. Risiko & bottleneck", "Kemungkinan/dampak: Rendah–Sedang–Tinggi",
                 ["ID", "Risiko", "Milestone terdampak", "Kemungkinan", "Dampak", "Mitigasi", "Pemilik", "Status"],
                 [r + ("PM", "Belum mulai") for r in RISKS], [6, 46, 16, 12, 10, 60, 8, 14], color=AMBER, status_col="H")
    simple_sheet(wb, "Step2_Outlook", "D. Post-SETAPAK 4 development outlook", "Bukan kewajiban capaian SETAPAK 4. Biaya = asumsi perencanaan",
                 ["Milestone", "Workstreams", "Syarat mulai", "Dokumen utama", "Durasi indikatif", "Pendanaan indikatif",
                  "Mitra potensial", "Ketergantungan pada STEP 1", "Batas waktu standar"],
                 [(s["code"], s["workstreams"], s["prereq"], s["docs"], s["duration"], s["funding"], s["partners"],
                   s["dep"], s["deadline"]) for s in STEP2],
                 [9, 50, 32, 32, 26, 20, 32, 32, 28], color=NAVY)
    simple_sheet(wb, "Transition_Checklist", "Output 5. Transition readiness checklist (pipeline → PDD)", "",
                 ["ID", "Kategori", "Persyaratan", "Bukti dari", "Status", "Catatan"],
                 [t + ("Belum mulai", "") for t in TRANSITION], [6, 22, 70, 14, 14, 30], status_col="E")
    simple_sheet(wb, "Sources", "Referensi", "Diakses Oktober 2026", ["Dokumen", "URL"], SOURCES, [60, 110])
    wb.save(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "matrix.xlsx")
