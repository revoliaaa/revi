# -*- coding: utf-8 -*-
"""Workplan Hutan Desa Pasir: STEP 1 (mingguan, s.d. Desember 2026) dan STEP 2 penyusunan PDD (bulanan, 2027).

    python3 build_workplan.py ../Workplan_HD_Pasir_2026-2027.xlsx ../Workplan_1Slide_HD_Pasir.pptx
"""
import datetime as dt
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pptx import Presentation
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from build_pptx import (AMBER, HAIR, INK, MUTED, NAVY, NAVY_MID, NAVY_TINT, TEAL, TEAL_MID, TEAL_TINT, WHITE, box,
                        line, rich, shape_text, text)

W1 = dt.date(2026, 10, 26)          # Senin minggu pertama STEP 1
WEEKS = 10                          # W1 (26 Okt) – W10 (28 Des)
PIN_SUBMIT_WEEK = 6                 # Jumat 4 Des 2026
BULAN = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]
MONTHS = 15                         # Jan 2027 – Mar 2028 (3 bulan terakhir = proses Plan Vivo setelah PDD diajukan)
TAF_END_MONTH = 9                   # Sep 2027


def week_label(w):
    a = W1 + dt.timedelta(weeks=w - 1)
    return f"W{w}\n{a.day} {BULAN[a.month - 1]}"


def month_label(m):
    return f"{BULAN[(m - 1) % 12]}\n{2027 + (m - 1) // 12}"


def wdate(w, end=False):
    return W1 + dt.timedelta(weeks=w - 1, days=4 if end else 0)


def mdate(m, end=False):
    y, mo = 2027 + (m - 1) // 12, (m - 1) % 12 + 1
    if not end:
        return dt.date(y, mo, 1)
    nxt = dt.date(y + (mo == 12), mo % 12 + 1, 1)
    return nxt - dt.timedelta(days=1)


# ---------------------------------------------------------------- STEP 1 (mingguan)
# (kode, tahap, kegiatan, output, PIC, mulai, selesai, milestone M, jenis)
# jenis: K = kegiatan, MS = milestone internal, EXT = keputusan pihak luar (indikatif)
STEP1 = [
    ("A", 1, "Tahap 1 — Persiapan Pendaftaran PIN", None, None, 1, 5, "M1–M5", "H"),
    ("A1", 1, "Kick-off LPHD–Pemdes–BPD; kesepakatan pendampingan & PKS SAMPAN–LPHD",
     "BA kick-off, surat kesediaan, PKS", "PM / Ketua LPHD", 1, 1, "M1", "K"),
    ("A2", 1, "Inventarisasi progres & dokumen yang sudah ada; susun data room",
     "Matriks progres + data room", "Field Facilitator", 1, 2, "M1.5", "K"),
    ("A3", 1, "Cek legal: SK Hutan Desa, SK LPHD, peta areal kerja, posisi hak karbon",
     "Memo legal awal", "Legal Officer", 1, 2, "M1.4", "K"),
    ("A4", 1, "Gap analysis data terhadap PIN Template V1.4", "Daftar data gap & rencana isi", "Tenaga ahli karbon", 2, 2,
     "M6.1", "K"),
    ("A5", 1, "Stakeholder analysis & konsultasi awal komunitas (Tool 1, 12)", "Stakeholder register, notulen",
     "Safeguards Officer", 2, 3, "M2", "K"),
    ("A6", 1, "Pengambilan data baseline dengan PV tools: problem analysis, resource & social mapping, FGD "
              "perempuan & kelompok rentan, historical timeline (Tool 2, 7, 9, 10, 13, 17)",
     "Output tools + peta partisipatif", "Field Facilitator / GIS", 2, 4, "M3.2–M3.3", "K"),
    ("A7", 1, "Analisis tutupan lahan & driver deforestasi/degradasi (citra 10–15 th)", "Peta tutupan & analisis driver",
     "GIS & Carbon Officer", 2, 4, "M5.2–M5.3", "K"),
    ("A8", 1, "Safeguard: E&S risk screening partisipatif; mekanisme pengaduan (GRM) interim",
     "Matriks risiko E&S, SOP GRM", "Safeguards Officer", 3, 4, "M3.4, M3.7", "K"),
    ("A9", 1, "Baseline scenario vs project scenario, intervensi & ToC, estimasi indikatif, keputusan metodologi",
     "Narasi scenario, memo metodologi, Go/No-Go", "Tenaga ahli karbon", 3, 5, "M5.4–M5.7", "K"),
    ("A10", 1, "FPIC: sesi informasi (W3) dan musyawarah keputusan tahap i–ii (W5) — jadwal ditentukan komunitas",
     "BA keputusan FPIC tahap i–ii", "Safeguards Officer / LPHD", 3, 5, "M3.1, M3.5", "K"),
    ("B", 2, "Tahap 2 — Submit PIN hingga Approval", None, None, 2, 10, "M6–M7", "H"),
    ("B1", 2, "Pre-submission call dengan Plan Vivo (eligibility & timeline)", "Catatan call", "Program Manager", 2, 2,
     "M6.4", "K"),
    ("B2", 2, "Penyusunan draft PIN (PIN Template V1.4)", "Draft PIN", "Tenaga ahli karbon", 4, 5, "M6.2", "K"),
    ("B3", 2, "Review PIN bersama LPHD (bahasa Indonesia) & finalisasi lampiran", "BA review, PIN final",
     "Field Facilitator / LPHD", 5, 6, "M6.3, M6.5", "K"),
    ("B4", 2, "Submit PIN + bayar PIN review fee (Jumat 4 Des 2026)", "Bukti email & kwitansi", "Direktur SAMPAN", 6, 6,
     "M6.6", "MS"),
    ("B5", 2, "Review PIN oleh Plan Vivo (±6 minggu; hasil ±pertengahan Jan 2027)", "PIN review report",
     "Plan Vivo", 7, 10, "M7.1", "EXT"),
    ("C", 3, "Tahap 3 — Masuk Pipeline PV (output utama)", None, None, 7, 10, "M7.5, M8", "H"),
    ("C1", 3, "Siapkan tim respons CAR/NIR (target ≤2 minggu)", "Tim & matriks respons", "Program Manager", 7, 8,
     "M7.2", "K"),
    ("C2", 3, "Profil proyek & teaser investor; daftar calon investor/offtaker", "Teaser, long list investor",
     "Program Manager", 7, 9, "M8", "K"),
    ("C3", 3, "Mulai registrasi pendamping/mitra teregistrasi (Permenhut 6/2026)", "Berkas registrasi",
     "Legal Officer", 7, 9, "M8", "K"),
    ("C4", 3, "Laporan akhir STEP 1 & paket handover ke penyusunan PDD", "Laporan + BA serah terima", "Program Manager",
     9, 10, "M7.6", "MS"),
]

# ---------------------------------------------------------------- STEP 2 (bulanan, 2027 → Q1 2028)
STEP2 = [
    ("D", 3, "Tindak lanjut review PIN → konfirmasi pipeline", None, None, 1, 2, "M7", "H"),
    ("D1", 3, "Tanggapi CAR/NIR PIN review report; revisi PIN (maks. 3 putaran)", "Matriks tanggapan, PIN revisi",
     "Tenaga ahli karbon", 1, 2, "M7.2–M7.3", "K"),
    ("D2", 3, "Konfirmasi PIN approved & tercantum di pipeline Plan Vivo (indikatif)", "Email approval, listing",
     "Plan Vivo", 2, 2, "M7.5", "EXT"),
    ("E", 4, "Persiapan penyusunan PDD", None, None, 1, 3, "M9", "H"),
    ("E1", 4, "Bentuk tim PDD; kontrak tenaga ahli karbon blue carbon; rencana kerja rinci", "SK tim, kontrak, workplan",
     "Direktur SAMPAN", 1, 1, "M9", "K"),
    ("E2", 4, "Konfirmasi metodologi: PM001 + modul PU001–PU005 dan status Mangrove Conservation Module",
     "Memo metodologi final (decision gate)", "Tenaga ahli karbon", 1, 2, "M9", "MS"),
    ("E3", 4, "Outline PDD (PDD Template V5.7) & pembagian penulis bab", "Outline PDD", "Tenaga ahli karbon", 2, 3, "M9",
     "K"),
    ("F", 4, "Data lapangan & baseline", None, None, 2, 8, "M9", "H"),
    ("F1", 4, "Baseline sosial-ekonomi rumah tangga (well-being, Tool 5) & GESI", "Dataset & laporan baseline sosial",
     "Safeguards Officer", 2, 4, "M9", "K"),
    ("F2", 4, "Stratifikasi & desain plot; inventarisasi biomassa & karbon tanah mangrove", "Data plot, laporan inventarisasi",
     "GIS & Carbon Officer / mitra riset", 3, 6, "M9", "K"),
    ("F3", 4, "Baseline biodiversitas", "Laporan biodiversitas", "Mitra riset", 4, 6, "M9", "K"),
    ("F4", 4, "Carbon accounting: baseline, ex-ante, leakage, uncertainty, risk buffer, additionality",
     "Spreadsheet & bab perhitungan karbon", "Tenaga ahli karbon", 6, 8, "M9", "K"),
    ("G", 4, "Safeguards, FPIC tahap iii & tata kelola", None, None, 3, 9, "M9", "H"),
    ("G1", 4, "E&S assessment penuh & rencana mitigasi", "Laporan E&S", "Safeguards Officer", 3, 5, "M9", "K"),
    ("G2", 4, "Benefit-sharing mechanism: draft → konsultasi → final (≥60% ke komunitas)", "Dokumen BSM disepakati",
     "Program Manager / LPHD", 5, 8, "M9", "K"),
    ("G3", 4, "Monitoring plan (karbon, sosial, biodiversitas) & peran WFD/KUPS", "Monitoring plan", "Tenaga ahli karbon",
     6, 8, "M9", "K"),
    ("G4", 4, "FPIC tahap iii (implementasi) & penandatanganan Project Agreement", "BA FPIC iii, Project Agreement",
     "Safeguards Officer / LPHD", 7, 9, "M9", "MS"),
    ("H", 4, "Penyusunan, konsultasi & pengajuan PDD", None, None, 5, 15, "M9", "H"),
    ("H1", 4, "Drafting bab PDD", "Draft PDD v0.5 → v0.9", "Tim PDD", 5, 9, "M9", "K"),
    ("H2", 4, "QA internal & review LPHD (bahasa Indonesia)", "BA review, PDD final", "Program Manager / LPHD", 9, 10,
     "M9", "K"),
    ("H3", 4, "Submit PDD ke Plan Vivo + fee (proyek harus sudah di pipeline)", "PDD submitted", "Direktur SAMPAN", 10, 10,
     "M9", "MS"),
    ("H4", 4, "PDD screening (±6 mgg) → konsultasi publik 4 mgg → PDD review (±8 mgg)", "Screening & review report",
     "Plan Vivo", 11, 15, "M9", "EXT"),
    ("I", 4, "Jalur nasional & investasi (paralel)", None, None, 1, 12, "M8, M10", "H"),
    ("I1", 3, "Registrasi pendamping/mitra teregistrasi selesai", "Bukti registrasi", "Legal Officer", 1, 3, "M8", "K"),
    ("I2", 3, "Investor/offtaker: penjajakan → LoI/term sheet → pre-financing validasi", "LoI / term sheet",
     "Program Manager", 2, 9, "M8", "K"),
    ("I3", 4, "Susun DPP & persiapan pencatatan ke Kemenhut (wajib sebelum validasi)", "Draft DPP", "Legal Officer",
     9, 12, "M10", "K"),
]

ASSUMPTIONS = [
    ("Target STEP 1", "Semua pekerjaan Tahap 1–2 selesai dan PIN terkirim Jumat 4 Desember 2026. Desember 2026 = PIN "
                      "dalam review Plan Vivo; paket handover siap."),
    ("Penerimaan pipeline", "Review awal ±6 minggu (Procedures Manual V3.9 §9) → hasil ±pertengahan Januari 2027; "
                            "dengan satu putaran CAR/NIR, konfirmasi pipeline ±Februari 2027. Agar diterima sebelum "
                            "akhir Desember, PIN harus terkirim ≤13 November 2026 tanpa CAR/NIR — hanya mungkin bila "
                            "data dan FPIC sudah tersedia."),
    ("FPIC", "Jadwal 5 minggu sangat ketat. Timeline FPIC ditetapkan pemegang hak (Project Requirements V5.7 §2.6.3). "
             "Bila komunitas butuh waktu lebih, tanggal submit PIN yang mundur."),
    ("Progres yang sudah ada", "Workplan mengasumsikan sebagian data, kelembagaan (LPHD, KUPS, WFD) dan dokumen legal "
                               "sudah tersedia dari kegiatan sebelumnya; A2 & A4 memastikan hal ini di 2 minggu pertama."),
    ("STEP 2 (2027)", "Fokus penyusunan PDD. Target: PDD diajukan Oktober 2027; screening, konsultasi publik dan review "
                      "Plan Vivo berjalan Nov 2027 – Mar 2028. PDD hanya bisa diajukan setelah proyek di pipeline."),
    ("Metodologi", "PDD wajib memakai metodologi yang sudah disetujui. Bila Mangrove Conservation Module belum "
                   "disetujui per Februari 2027, ruang lingkup PDD difokuskan ke komponen restorasi (PM001)."),
    ("Pendanaan", "TAF s.d. September 2027 (dikonfirmasi apakah mencakup kegiatan PDD). Okt–Des 2027 dan biaya "
                  "fee/validasi dari pre-financing investor atau pendanaan lanjutan."),
    ("Libur akhir tahun", "W9–W10 (21 Des – 1 Jan) diasumsikan minim kegiatan lapangan; review Plan Vivo bisa melambat."),
]


# ---------------------------------------------------------------- Excel
TEALH, NAVYH, AMBERH = "0E7C6B", "1F2A44", "C26A0A"
thin = Side(style="thin", color="D5DBD6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
STATUS = ["Belum mulai", "Berjalan", "Selesai", "Tertunda"]


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def gantt_sheet(wb, name, title, sub, rows, ncols, col_label, start_date, end_date, color, ext_after=None,
                marker_col=None):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color=color)
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, size=9, color="5F6B65")
    head = ["Kode", "Tahap", "Kegiatan", "Output", "PIC", "Milestone", "Mulai", "Selesai"]
    head += [col_label(i) for i in range(1, ncols + 1)] + ["Status", "Catatan"]
    for i, h in enumerate(head, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.fill = fill(color)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER
    if ext_after:
        for i in range(ext_after + 1, ncols + 1):
            ws.cell(row=4, column=8 + i).fill = fill(AMBERH)
    ws.row_dimensions[4].height = 30
    r = 5
    for code, tahap, keg, out, pic, a, z, ms, kind in rows:
        head_row = kind == "H"
        vals = [code, f"Tahap {tahap}", keg, out or "", pic or "", ms, start_date(a), end_date(z, True)]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
            if i in (7, 8):
                c.number_format = "dd mmm yyyy"
            if head_row:
                c.font = Font(bold=True, color=color)
                c.fill = fill("E3F1EE" if color == TEALH else "E6E9EF")
        for i in range(1, ncols + 1):
            c = ws.cell(row=r, column=8 + i)
            c.border = BORDER
            if a <= i <= z:
                if head_row:
                    c.fill = fill("B9D9D2" if color == TEALH else "BFC6D4")
                elif kind == "EXT":
                    c.fill = fill(AMBERH)
                    c.value = "◇" if i == a else None
                elif kind == "MS":
                    c.fill = fill(color)
                    c.value = "◆" if i == z else None
                    c.font = Font(color="FFFFFF", bold=True)
                    c.alignment = Alignment(horizontal="center")
                else:
                    c.fill = fill("7FBCB1" if color == TEALH else "8A93A8")
        if not head_row:
            ws.cell(row=r, column=9 + ncols, value="Belum mulai").border = BORDER
        ws.cell(row=r, column=10 + ncols).border = BORDER
        r += 1
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    sc = get_column_letter(9 + ncols)
    dv.add(f"{sc}5:{sc}{r - 1}")
    ws.conditional_formatting.add(f"{sc}5:{sc}{r - 1}",
                                  CellIsRule(operator="equal", formula=['"Selesai"'], fill=fill("E3F1EE")))
    ws.conditional_formatting.add(f"{sc}5:{sc}{r - 1}",
                                  CellIsRule(operator="equal", formula=['"Tertunda"'], fill=fill("F6EBDD")))
    widths = [6, 8, 52, 30, 20, 11, 11, 11] + [6.5] * ncols + [13, 26]
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "D5"
    r += 1
    ws.cell(row=r, column=3, value="Legenda: blok berwarna = kegiatan · ◆ = milestone internal · ◇ oranye = "
                                   "keputusan pihak luar (indikatif) · header oranye = di luar periode TAF/program").font = \
        Font(italic=True, size=9, color="5F6B65")
    return ws


def build_xlsx(out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws["A1"] = "Workplan Community Carbon Hutan Desa Pasir — Plan Vivo PV Climate"
    ws["A1"].font = Font(bold=True, size=14, color=TEALH)
    ws["A2"] = "STEP 1 (Tahap 1–3) s.d. Desember 2026 · STEP 2 penyusunan PDD (Tahap 4) sepanjang 2027"
    ws["A2"].font = Font(italic=True, size=9, color="5F6B65")
    rows = [("Periode", "Fokus", "Milestone kunci", "Target")] + [
        ("STEP 1 · 26 Okt – 31 Des 2026", "Tahap 1 Persiapan PIN → Tahap 2 Submit PIN → Tahap 3 masuk pipeline",
         "PIN terkirim + fee dibayar", "Jumat 4 Des 2026"),
        ("", "", "Paket handover ke PDD", "Akhir Des 2026"),
        ("", "", "PIN approved, proyek di pipeline (keputusan Plan Vivo)", "Indikatif Jan–Feb 2027"),
        ("STEP 2 · Jan – Des 2027", "Tahap 4 Penyusunan PDD (+ investasi & jalur nasional paralel)",
         "Memo metodologi final", "Feb 2027"),
        ("", "", "Data lapangan karbon & sosial lengkap", "Jun 2027"),
        ("", "", "FPIC tahap iii & Project Agreement", "Sep 2027"),
        ("", "", "PDD diajukan ke Plan Vivo", "Okt 2027"),
        ("", "", "PDD screening → konsultasi publik → review (Plan Vivo)", "Indikatif Nov 2027 – Mar 2028"),
    ]
    for i, row in enumerate(rows, 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
            if i == 4:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = fill(NAVYH)
    r = 4 + len(rows) + 1
    ws.cell(row=r, column=1, value="Asumsi & catatan").font = Font(bold=True, color=TEALH)
    for k, v in ASSUMPTIONS:
        r += 1
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=r, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 45
    for i, wd in enumerate([30, 52, 46, 26], 1):
        ws.column_dimensions[get_column_letter(i)].width = wd

    gantt_sheet(wb, "Step1_Mingguan", "STEP 1 — Workplan mingguan (26 Okt – 31 Des 2026)",
                "Target: PIN terkirim 4 Des 2026 (W6). Keputusan Plan Vivo indikatif Jan–Feb 2027.",
                STEP1, WEEKS, lambda i: week_label(i).replace("\n", " "), wdate, wdate, TEALH)
    gantt_sheet(wb, "Step2_PDD_2027", "STEP 2 — Workplan penyusunan PDD (Jan 2027 – Mar 2028)",
                "Target: PDD diajukan Okt 2027. Header oranye = setelah TAF (Sep 2027) / proses Plan Vivo.",
                STEP2, MONTHS, lambda i: month_label(i).replace("\n", " "), mdate, mdate, NAVYH,
                ext_after=TAF_END_MONTH)
    wb.save(out)


# ---------------------------------------------------------------- Slide
def build_slide(out):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    from pptx.dml.color import RGBColor
    box(s, 0, 0, 13.333, 7.5, fill=RGBColor(0xFB, 0xFB, 0xF9))
    text(s, 0.45, 0.3, 11, 0.25, "WORKPLAN · COMMUNITY CARBON HUTAN DESA PASIR · PLAN VIVO PV CLIMATE", size=9, bold=True,
         color=TEAL)
    text(s, 0.45, 0.55, 12.4, 0.8, "PIN terkirim awal Desember 2026; sepanjang 2027 tim menyusun PDD\n"
                                   "hingga diajukan ke Plan Vivo pada Oktober 2027", size=20, color=INK)

    lab_w = 3.3
    xL = 0.45 + lab_w                   # awal gantt
    xM = xL + 10 * 0.33                 # pemisah minggu | bulan
    xR = 12.93
    ww = (xM - xL) / WEEKS
    mw = (xR - xM - 0.12) / 12          # tampilkan Jan–Des 2027
    xM2 = xM + 0.12
    y0 = 1.62

    # header zona
    hb = box(s, xL, y0, xM - xL, 0.3, fill=TEAL)
    shape_text(hb, "STEP 1 · Okt – Des 2026 (mingguan)", size=8)
    hb = box(s, xM2, y0, xR - xM2, 0.3, fill=NAVY)
    shape_text(hb, "STEP 2 · Penyusunan PDD 2027 (bulanan)", size=8)
    for w in range(1, WEEKS + 1):
        a = W1 + dt.timedelta(weeks=w - 1)
        text(s, xL + (w - 1) * ww, y0 + 0.33, ww, 0.3, f"{a.day}\n{BULAN[a.month - 1]}", size=6, color=MUTED,
             align=PP_ALIGN.CENTER)
    for m in range(1, 13):
        text(s, xM2 + (m - 1) * mw, y0 + 0.36, mw, 0.2, BULAN[m - 1], size=7, color=MUTED, align=PP_ALIGN.CENTER)

    # baris: workstream ringkas (kode, label, segmen-segmen) ; segmen = (zona, a, z, jenis)
    rows = [
        ("T1", "Inventarisasi progres, legal & gap data", [("w", 1, 2, "K")]),
        ("T1", "Baseline dengan PV tools & analisis citra", [("w", 2, 4, "K")]),
        ("T1", "Safeguard: E&S screening, GRM, FPIC i–ii", [("w", 3, 5, "K")]),
        ("T1", "Baseline & project scenario, metodologi", [("w", 3, 5, "K")]),
        ("T2", "Draft PIN → review LPHD → submit", [("w", 4, 6, "K"), ("w", 6, 6, "MS")]),
        ("T2", "Review Plan Vivo → CAR/NIR", [("w", 7, 10, "EXT"), ("m", 1, 2, "EXT")]),
        ("T3", "Masuk pipeline · teaser investor", [("w", 7, 10, "K"), ("m", 2, 2, "EXT")]),
        ("T4", "Tim PDD & keputusan metodologi", [("m", 1, 2, "K"), ("m", 2, 2, "MS")]),
        ("T4", "Data lapangan: karbon, sosial, biodiversitas", [("m", 2, 6, "K")]),
        ("T4", "Carbon accounting & monitoring plan", [("m", 6, 8, "K")]),
        ("T4", "Benefit-sharing & FPIC iii, Project Agreement", [("m", 5, 9, "K"), ("m", 9, 9, "MS")]),
        ("T4", "Drafting, QA & review LPHD → submit PDD", [("m", 5, 10, "K"), ("m", 10, 10, "MS")]),
        ("T4", "Screening & konsultasi publik Plan Vivo", [("m", 11, 12, "EXT")]),
        ("—", "Paralel: pendamping teregistrasi, investor, DPP", [("m", 1, 12, "P")]),
    ]
    ry = y0 + 0.66
    rh = (6.25 - ry) / len(rows)
    for i, (tag, lbl, segs) in enumerate(rows):
        y = ry + i * rh
        col = TEAL if tag in ("T1", "T2", "T3") else NAVY
        if i % 2 == 0:
            box(s, 0.45, y, xR - 0.45, rh, fill=RGBColor(0xF2, 0xF4, 0xF1))
        tb = box(s, 0.5, y + rh * 0.18, 0.42, rh * 0.64, fill=col if tag != "—" else MUTED,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        shape_text(tb, tag.replace("T", "Tahap ") if False else tag, size=7)
        text(s, 1.0, y, lab_w - 0.6, rh, lbl, size=8, color=INK, anchor=MSO_ANCHOR.MIDDLE)
        for zone, a, z, kind in segs:
            if zone == "w":
                x, wd = xL + (a - 1) * ww, (z - a + 1) * ww
            else:
                x, wd = xM2 + (a - 1) * mw, (z - a + 1) * mw
            if kind == "MS":
                cxm = x + wd * (0.8 if zone == "w" else 0.5)
                d = box(s, cxm - 0.11, y + rh / 2 - 0.11, 0.22, 0.22, fill=col, line=WHITE, shape=MSO_SHAPE.DIAMOND,
                        lw=1)
                continue
            fillc = {"K": TEAL_MID if col == TEAL else NAVY_MID, "EXT": AMBER, "P": NAVY_TINT}[kind]
            box(s, x + 0.02, y + rh * 0.28, wd - 0.04, rh * 0.44, fill=fillc,
                line=NAVY_MID if kind == "P" else None, lw=0.5)

    # garis penanda
    xs = xL + (PIN_SUBMIT_WEEK - 0.2) * ww
    line(s, xs, y0 + 0.64, xs, 6.25, color=TEAL, w=1.25, dash=MSO_LINE.DASH)
    text(s, xs - 0.9, 6.27, 1.8, 0.2, "PIN submit · 4 Des", size=7, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    line(s, xM + 0.06, y0, xM + 0.06, 6.25, color=HAIR, w=1)
    xt = xM2 + TAF_END_MONTH * mw
    line(s, xt, y0 + 0.64, xt, 6.25, color=AMBER, w=1.25, dash=MSO_LINE.DASH)
    text(s, xt - 1.85, 6.27, 1.8, 0.2, "Akhir TAF · Sep 2027", size=7, bold=True, color=AMBER, align=PP_ALIGN.RIGHT)
    text(s, xt + 0.05, 6.27, 2.0, 0.2, "PDD submit · Okt 2027", size=7, bold=True, color=NAVY)

    # strip bawah
    ky = 6.6
    line(s, 0.45, ky - 0.06, 12.93, ky - 0.06, color=HAIR, w=0.75)
    lg = [(TEAL_MID, "Kegiatan STEP 1"), (NAVY_MID, "Kegiatan STEP 2"), (AMBER, "Keputusan Plan Vivo (indikatif)")]
    for lx, (c, l) in zip((0.45, 1.95, 3.45), lg):
        box(s, lx, ky + 0.07, 0.3, 0.12, fill=c)
        text(s, lx + 0.35, ky, 2.2, 0.25, l, size=8, color=MUTED)
    lx = 5.75
    d = box(s, lx, ky + 0.02, 0.2, 0.2, fill=MUTED, shape=MSO_SHAPE.DIAMOND)
    text(s, lx + 0.27, ky, 1.6, 0.25, "Milestone", size=8, color=MUTED)
    rich(s, 7.15, ky - 0.03, 5.8, 0.6, [[("Catatan: ", True, AMBER),
                                          ("review PIN ±6 minggu → keputusan pipeline indikatif Jan–Feb 2027. PDD "
                                           "hanya dapat diajukan setelah proyek di pipeline dan memakai metodologi "
                                           "yang sudah disetujui.", False, INK)]], size=7.5)
    text(s, 0.45, 7.15, 12.4, 0.22, "Sumber: Plan Vivo PV Climate Procedures Manual V3.9 (§2.2, §4, §9); Project "
                                    "Requirements V5.7; Permenhut 6/2026.", size=7, color=MUTED)
    s.notes_slide.notes_text_frame.text = (
        "STEP 1 berjalan mingguan dari 26 Oktober sampai akhir Desember 2026, dengan PIN terkirim 4 Desember. Keputusan "
        "Plan Vivo atas PIN diperkirakan Januari–Februari 2027. STEP 2 sepanjang 2027 fokus pada penyusunan PDD: "
        "metodologi dikunci Februari, data lapangan Februari–Juni, perhitungan karbon dan monitoring plan Juni–Agustus, "
        "benefit-sharing dan FPIC tahap iii sampai September, PDD diajukan Oktober 2027.")
    prs.save(out)


if __name__ == "__main__":
    build_xlsx(sys.argv[1] if len(sys.argv) > 1 else "workplan.xlsx")
    build_slide(sys.argv[2] if len(sys.argv) > 2 else "workplan.pptx")
